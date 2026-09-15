"""
Lightweight document text extractors for attachments.

OCR default provider: pytesseract / Tesseract 5 (print fast-path).
R5.3.2 / R5.5A: low_quality / unusable / empty → RapidOCR PP-OCRv6 small
in a short-lived subprocess. PaddleOCR remains an unused extension hook.
"""
from __future__ import annotations

import tempfile
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.services.ocr_quality import (
    OCR_QUALITY_EMPTY,
    OCR_QUALITY_LOW,
    OCR_QUALITY_OK,
    OCR_QUALITY_UNUSABLE,
    OCR_QUALITY_UNKNOWN,
    OcrQualityAssessment,
    PROFILE_UNKNOWN,
    THRESHOLDS,
    assess_ocr_text,
    merge_quality_into_metadata,
    provider_stability_rank,
)
from app.services.ocr_runtime import (
    OCR_DISABLED,
    OCR_FAILED,
    OCR_STATUS_DISABLED,
    OCR_STATUS_EMPTY,
    OCR_STATUS_FAILED,
    OCR_STATUS_SUCCEEDED,
    OCR_STATUS_UNAVAILABLE,
    OcrRequiredError,
    PdfScanOcrUnavailableError,
    check_ocr_runtime,
    normalize_provider,
)


def chunk_text(text: str, *, size: int = 800, overlap: int = 100) -> list[str]:
    """Split text into overlapping chunks."""
    cleaned = "\n".join(line.strip() for line in (text or "").splitlines() if line.strip())
    if not cleaned:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + size)
        chunk = cleaned[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
    return chunks


def _settings():
    from app.config import settings

    return settings


def _pdf_ocr_fallback_enabled() -> bool:
    return bool(_settings().attachment_pdf_ocr_fallback_enabled)


def _pdf_ocr_dpi() -> int:
    return max(72, min(int(_settings().attachment_pdf_ocr_dpi or 180), 400))


def _ocr_language_label() -> str:
    prov = normalize_provider()
    if prov == "paddleocr":
        return str(_settings().attachment_paddleocr_lang or "ch")
    return str(_settings().attachment_tesseract_lang or "chi_sim+eng")


def _render_pdf_page_to_temp_image(page: Any) -> Path:
    dpi = _pdf_ocr_dpi()
    pixmap = page.get_pixmap(dpi=dpi)
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    tmp_path = Path(tmp.name)
    tmp.close()
    try:
        pixmap.save(str(tmp_path))
    except TypeError:
        pixmap.save(tmp_path)
    return tmp_path


def _extract_pdf_page_ocr_chunks(page: Any, page_no: int) -> list[dict]:
    """OCR one textless PDF page.

    Unavailable/failed/disabled OR final non-searchable OCR → fail whole attachment
    (no silent partial index for mixed PDFs).
    """
    image_path = _render_pdf_page_to_temp_image(page)
    try:
        ocr_text, provider, status, error_code, safe_meta = _run_ocr(image_path)
    finally:
        try:
            image_path.unlink(missing_ok=True)
        except OSError:
            pass

    if status in {
        OCR_STATUS_UNAVAILABLE,
        OCR_STATUS_FAILED,
        OCR_STATUS_DISABLED,
    }:
        # Do not return partial digital-page chunks for a mixed PDF.
        raise PdfScanOcrUnavailableError(error_code=error_code or OCR_FAILED)

    quality = assess_ocr_text(
        ocr_text or "",
        runtime_status=status,
        mean_confidence=safe_meta.get("ocr_mean_confidence"),
        provider=provider,
        provider_version=str(safe_meta.get("ocr_model") or provider),
    )

    # Final OCR must be searchable; empty/low/unusable must not silently drop a scan page.
    if not quality.search_eligible:
        raise PdfScanOcrUnavailableError(error_code=error_code or OCR_FAILED)

    if status != OCR_STATUS_SUCCEEDED or not ocr_text:
        raise PdfScanOcrUnavailableError(error_code=error_code or OCR_FAILED)

    lang = _ocr_language_label()
    chunks = []
    for chunk in chunk_text(ocr_text):
        meta = {
            "extraction_method": "ocr",
            "source": "pymupdf_page_render",
            "ocr_provider": provider,
            "ocr_language": lang,
            "ocr_status": OCR_STATUS_SUCCEEDED,
            "ocr_model": safe_meta.get("ocr_model"),
            "ocr_fallback_used": bool(safe_meta.get("ocr_fallback_used")),
            "ocr_elapsed_ms": safe_meta.get("ocr_elapsed_ms"),
            "ocr_attempts": safe_meta.get("ocr_attempts"),
        }
        chunks.append({
            "modality": "ocr",
            "page_no": page_no,
            "slide_no": None,
            "content": chunk,
            "metadata_json": merge_quality_into_metadata(meta, quality),
        })
    return chunks


def extract_pdf_chunks(path: Path) -> list[dict]:
    """Extract text chunks from a digital PDF; OCR only when text layer empty + fallback on.

    When fallback is enabled and any textless page hits unavailable/failed OCR,
    raise PdfScanOcrUnavailableError so the attachment becomes failed (no partial indexed).
    """
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PDF 文本提取组件未就绪") from exc

    extracted: list[dict] = []
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            text_chunks = chunk_text(text)
            for chunk in text_chunks:
                extracted.append({
                    "modality": "pdf_text",
                    "page_no": page_index,
                    "slide_no": None,
                    "content": chunk,
                    "metadata_json": {
                        "extraction_method": "text_layer",
                        "source": "pymupdf",
                    },
                })
            # Never OCR pages that already have a usable text layer.
            if not text_chunks and _pdf_ocr_fallback_enabled():
                extracted.extend(_extract_pdf_page_ocr_chunks(page, page_index))
    return extracted


def extract_pptx_chunks(path: Path) -> list[dict]:
    """Extract text chunks from a PPTX file."""
    try:
        from pptx import Presentation  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PPTX 文本提取组件未就绪") from exc

    prs = Presentation(str(path))
    extracted: list[dict] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        text_parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    run.text
                    for paragraph in shape.text_frame.paragraphs
                    for run in paragraph.runs
                    if run.text
                )
                if text.strip():
                    text_parts.append(text.strip())
        for chunk in chunk_text("\n".join(text_parts)):
            extracted.append({
                "modality": "ppt_text",
                "page_no": None,
                "slide_no": slide_index,
                "content": chunk,
                "metadata_json": {
                    "extraction_method": "presentation_text",
                    "source": "python-pptx",
                },
            })
    return extracted


def _collect_paddle_text(value: Any) -> list[str]:
    texts: list[str] = []
    if isinstance(value, (list, tuple)):
        if (
            len(value) >= 2
            and isinstance(value[1], (list, tuple))
            and value[1]
            and isinstance(value[1][0], str)
        ):
            texts.append(value[1][0])
            return texts
        for item in value:
            texts.extend(_collect_paddle_text(item))
    return texts


def _ocr_with_pytesseract(path: Path) -> tuple[str, Optional[float]]:
    """Return (text, mean_confidence in 0..1 or None if unavailable).

    Caller should pass an OCR-normalized RGB PNG path (see ocr_image_normalize).
    """
    import pytesseract  # type: ignore
    from PIL import Image  # type: ignore

    lang = str(_settings().attachment_tesseract_lang or "chi_sim+eng")
    cmd = (_settings().attachment_tesseract_cmd or "").strip()
    previous_cmd = None
    if cmd:
        previous_cmd = getattr(pytesseract.pytesseract, "tesseract_cmd", None)
        pytesseract.pytesseract.tesseract_cmd = cmd
    try:
        with Image.open(path) as image:
            # Reject raw multi-frame containers that pytesseract cannot consume.
            fmt = str(getattr(image, "format", "") or "").upper()
            n_frames = int(getattr(image, "n_frames", 1) or 1)
            if fmt == "MPO" or n_frames > 1:
                raise RuntimeError("OCR_IMAGE_NOT_NORMALIZED")
            text = pytesseract.image_to_string(image, lang=lang).strip()
            mean_conf: Optional[float] = None
            try:
                data = pytesseract.image_to_data(
                    image, lang=lang, output_type=pytesseract.Output.DICT
                )
                confs: list[float] = []
                for raw in data.get("conf") or []:
                    try:
                        value = float(raw)
                    except Exception:
                        continue
                    if value >= 0:
                        confs.append(value)
                if confs:
                    mean_conf = (sum(confs) / len(confs)) / 100.0
            except Exception:
                mean_conf = None
            return text, mean_conf
    finally:
        if cmd and previous_cmd is not None:
            try:
                pytesseract.pytesseract.tesseract_cmd = previous_cmd
            except Exception:
                pass


def _ocr_with_paddleocr(path: Path) -> str:
    from paddleocr import PaddleOCR  # type: ignore

    lang = str(_settings().attachment_paddleocr_lang or "ch")
    use_angle_cls = bool(_settings().attachment_paddleocr_use_angle_cls)
    ocr = PaddleOCR(use_angle_cls=use_angle_cls, lang=lang)
    try:
        result = ocr.ocr(str(path), cls=use_angle_cls)
    except TypeError:
        result = ocr.ocr(str(path))
    return "\n".join(item for item in _collect_paddle_text(result) if item.strip()).strip()


def _quality_rank(assessment: OcrQualityAssessment) -> int:
    order = {
        OCR_QUALITY_OK: 4,
        OCR_QUALITY_LOW: 2,
        OCR_QUALITY_UNUSABLE: 1,
        OCR_QUALITY_EMPTY: 0,
        "unknown": 0,
    }
    return int(order.get(assessment.ocr_quality_status, 0))


def _needs_rapidocr_fallback(
    assessment: OcrQualityAssessment,
    status: str,
    *,
    mean_confidence: Optional[float] = None,
) -> bool:
    if status == OCR_STATUS_EMPTY:
        return True
    if assessment.ocr_quality_status in {
        OCR_QUALITY_LOW,
        OCR_QUALITY_UNUSABLE,
        OCR_QUALITY_EMPTY,
        OCR_QUALITY_UNKNOWN,
    }:
        return True
    # Uncertain / abnormal structure → verify with RapidOCR.
    if assessment.content_profile == PROFILE_UNKNOWN:
        return True
    if float(assessment.profile_confidence or 0.0) < THRESHOLDS.profile_confidence_floor:
        return True
    if float(assessment.quality_score or 0.0) < THRESHOLDS.quality_score_verify_below:
        return True
    # Primary "ok" with weak measured confidence: verify via RapidOCR and pick better.
    # Missing confidence keeps the R5.5A fast-path (synthetic / string-only OCR).
    if assessment.search_eligible and assessment.ocr_quality_status == OCR_QUALITY_OK:
        if mean_confidence is not None:
            threshold = float(
                getattr(_settings(), "attachment_rapidocr_verify_below_confidence", 0.85) or 0.85
            )
            if float(mean_confidence) < threshold:
                return True
    return False


def _safe_attempt_meta(
    *,
    provider: str,
    model: Optional[str],
    quality: str,
    elapsed_ms: float,
    error_code: Optional[str],
    mean_confidence: Optional[float],
) -> Dict[str, Any]:
    """Metadata-safe attempt row: no paths, body, or exception stacks."""
    row: Dict[str, Any] = {
        "provider": provider,
        "quality": quality,
        "elapsed_ms": round(float(elapsed_ms), 2),
    }
    if model:
        row["model"] = model
    if mean_confidence is not None:
        row["confidence"] = round(float(mean_confidence), 4)
    if error_code:
        row["error_code"] = str(error_code)
    return row


def _pick_better_ocr(
    primary: Dict[str, Any],
    fallback: Dict[str, Any],
) -> Dict[str, Any]:
    """
    Choose higher-quality text; never concatenate providers.

    Order: search_eligible → quality rank → quality_score → confidence →
    meaningful chars → stable provider rank. No generic cjk_ratio tie-break.
    """
    a: OcrQualityAssessment = primary["assessment"]
    b: OcrQualityAssessment = fallback["assessment"]

    if bool(b.search_eligible) != bool(a.search_eligible):
        return fallback if b.search_eligible else primary

    rb = _quality_rank(b)
    ra = _quality_rank(a)
    if rb != ra:
        return fallback if rb > ra else primary

    sb = float(getattr(b, "quality_score", 0.0) or 0.0)
    sa = float(getattr(a, "quality_score", 0.0) or 0.0)
    if abs(sb - sa) > 1e-9:
        return fallback if sb > sa else primary

    bc = b.ocr_mean_confidence if b.ocr_mean_confidence is not None else -1.0
    ac = a.ocr_mean_confidence if a.ocr_mean_confidence is not None else -1.0
    if bc != ac:
        return fallback if bc > ac else primary

    mb = int(getattr(b, "meaningful_char_count", 0) or 0)
    ma = int(getattr(a, "meaningful_char_count", 0) or 0)
    if mb != ma:
        return fallback if mb > ma else primary

    pb = provider_stability_rank(str(fallback.get("provider") or ""))
    pa = provider_stability_rank(str(primary.get("provider") or ""))
    if pb != pa:
        return fallback if pb < pa else primary
    return primary


def _run_primary_ocr(path: Path) -> Dict[str, Any]:
    """Run configured primary OCR (default Tesseract). No RapidOCR here."""
    from app.services.ocr_runtime import OCR_ENGINE_UNAVAILABLE

    provider = normalize_provider()
    started = time.perf_counter()
    if provider in {"disabled", "none", "off"}:
        assessment = assess_ocr_text(
            "",
            runtime_status=OCR_STATUS_DISABLED,
            provider="disabled",
            provider_version="disabled",
        )
        return {
            "text": "",
            "provider": "disabled",
            "model": None,
            "status": OCR_STATUS_DISABLED,
            "error_code": OCR_DISABLED,
            "assessment": assessment,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "mean_confidence": None,
        }

    # Primary chain remains tesseract even if provider string is rapidocr
    # (rapidocr is fallback-only unless explicitly forced later).
    primary = provider if provider not in {"rapidocr", "rapidocr_ppocrv6_small"} else "pytesseract"
    report = check_ocr_runtime(provider=primary)
    if primary in {"pytesseract", "tesseract"} and not report.get("ready"):
        assessment = assess_ocr_text(
            "",
            runtime_status=OCR_STATUS_UNAVAILABLE,
            provider="pytesseract",
            provider_version="pytesseract",
        )
        return {
            "text": "",
            "provider": "pytesseract",
            "model": "tesseract5",
            "status": OCR_STATUS_UNAVAILABLE,
            "error_code": str(report.get("error_code") or OCR_ENGINE_UNAVAILABLE),
            "assessment": assessment,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "mean_confidence": None,
        }

    if primary == "paddleocr" and not report.get("python_package", {}).get("available"):
        assessment = assess_ocr_text(
            "",
            runtime_status=OCR_STATUS_UNAVAILABLE,
            provider="paddleocr",
            provider_version="paddleocr",
        )
        return {
            "text": "",
            "provider": "paddleocr",
            "model": None,
            "status": OCR_STATUS_UNAVAILABLE,
            "error_code": str(report.get("error_code") or OCR_ENGINE_UNAVAILABLE),
            "assessment": assessment,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "mean_confidence": None,
        }

    mean_confidence: Optional[float] = None
    try:
        if primary in {"", "pytesseract", "tesseract"}:
            text, mean_confidence = _ocr_with_pytesseract(path)
            prov_name = "pytesseract"
            model = "tesseract5"
        elif primary == "paddleocr":
            text = _ocr_with_paddleocr(path)
            prov_name = "paddleocr"
            model = "paddleocr"
        else:
            assessment = assess_ocr_text(
                "",
                runtime_status=OCR_STATUS_UNAVAILABLE,
                provider=primary,
                provider_version=primary,
            )
            return {
                "text": "",
                "provider": primary,
                "model": None,
                "status": OCR_STATUS_UNAVAILABLE,
                "error_code": OCR_DISABLED,
                "assessment": assessment,
                "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
                "mean_confidence": None,
            }
    except ImportError:
        assessment = assess_ocr_text(
            "",
            runtime_status=OCR_STATUS_UNAVAILABLE,
            provider=primary or "pytesseract",
            provider_version=primary or "pytesseract",
        )
        return {
            "text": "",
            "provider": primary or "pytesseract",
            "model": None,
            "status": OCR_STATUS_UNAVAILABLE,
            "error_code": OCR_ENGINE_UNAVAILABLE,
            "assessment": assessment,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "mean_confidence": None,
        }
    except Exception:
        assessment = assess_ocr_text(
            "",
            runtime_status=OCR_STATUS_FAILED,
            provider=primary or "pytesseract",
            provider_version=primary or "pytesseract",
        )
        return {
            "text": "",
            "provider": primary or "pytesseract",
            "model": None,
            "status": OCR_STATUS_FAILED,
            "error_code": OCR_FAILED,
            "assessment": assessment,
            "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
            "mean_confidence": None,
        }

    status = OCR_STATUS_SUCCEEDED if text else OCR_STATUS_EMPTY
    assessment = assess_ocr_text(
        text or "",
        runtime_status=status,
        mean_confidence=mean_confidence,
        provider=prov_name,
        provider_version=prov_name,
    )
    return {
        "text": text or "",
        "provider": prov_name,
        "model": model,
        "status": status,
        "error_code": None,
        "assessment": assessment,
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "mean_confidence": assessment.ocr_mean_confidence,
    }


def _should_try_rapidocr_fallback(primary: Dict[str, Any]) -> bool:
    """Fallback on empty/low/unusable/failed; also when primary unavailable but RapidOCR ready."""
    from app.services.ocr_rapidocr import check_rapidocr_runtime, rapidocr_fallback_enabled

    if not rapidocr_fallback_enabled():
        return False
    status = str(primary.get("status") or "")
    if status == OCR_STATUS_DISABLED:
        return False
    if status == OCR_STATUS_FAILED:
        return True
    if status == OCR_STATUS_UNAVAILABLE:
        rapid = check_rapidocr_runtime()
        return bool(rapid.get("ready"))
    return _needs_rapidocr_fallback(
        primary["assessment"],
        status,
        mean_confidence=primary.get("mean_confidence"),
    )


def _run_ocr(path: Path) -> tuple[str, str, str, str | None, Dict[str, Any]]:
    """
    Execute OCR with Tesseract fast-path and optional RapidOCR quality fallback.

    Returns: (text, provider, ocr_status, error_code|None, safe_meta)
    Does not raise for optional OCR failures — caller decides required semantics.
    Embedding must only run after this returns (RapidOCR subprocess already exited).
    """
    from app.services.ocr_image_normalize import normalized_ocr_image
    from app.services.ocr_runtime import OCR_ENGINE_UNAVAILABLE

    provider = normalize_provider()
    if provider in {"disabled", "none", "off"}:
        safe_meta: Dict[str, Any] = {
            "ocr_provider": "disabled",
            "ocr_model": None,
            "ocr_fallback_used": False,
            "ocr_mean_confidence": None,
            "ocr_elapsed_ms": 0,
            "ocr_attempts": [],
            "ocr_error_code": OCR_DISABLED,
        }
        return "", "disabled", OCR_STATUS_DISABLED, OCR_DISABLED, safe_meta

    try:
        with normalized_ocr_image(path) as (ocr_path, norm_meta):
            return _run_ocr_on_normalized(ocr_path, norm_meta=norm_meta)
    except ImportError:
        provider = normalize_provider() or "pytesseract"
        safe_meta: Dict[str, Any] = {
            "ocr_provider": provider,
            "ocr_model": None,
            "ocr_fallback_used": False,
            "ocr_mean_confidence": None,
            "ocr_elapsed_ms": 0,
            "ocr_attempts": [],
            "ocr_error_code": OCR_ENGINE_UNAVAILABLE,
        }
        return "", provider, OCR_STATUS_UNAVAILABLE, OCR_ENGINE_UNAVAILABLE, safe_meta
    except Exception:
        provider = normalize_provider() or "pytesseract"
        safe_meta = {
            "ocr_provider": provider,
            "ocr_model": None,
            "ocr_fallback_used": False,
            "ocr_mean_confidence": None,
            "ocr_elapsed_ms": 0,
            "ocr_attempts": [],
            "ocr_error_code": OCR_FAILED,
        }
        return "", provider, OCR_STATUS_FAILED, OCR_FAILED, safe_meta


def _run_ocr_on_normalized(
    path: Path,
    *,
    norm_meta: Optional[Dict[str, Any]] = None,
) -> tuple[str, str, str, str | None, Dict[str, Any]]:
    from app.services.ocr_rapidocr import (
        PROVIDER_NAME as RAPID_PROVIDER,
        run_rapidocr_subprocess,
    )

    primary = _run_primary_ocr(path)
    attempts: List[Dict[str, Any]] = [
        _safe_attempt_meta(
            provider=str(primary["provider"]),
            model=primary.get("model"),
            quality=str(primary["assessment"].ocr_quality_status),
            elapsed_ms=float(primary["elapsed_ms"]),
            error_code=primary.get("error_code"),
            mean_confidence=primary.get("mean_confidence"),
        )
    ]
    chosen = primary
    fallback_used = False

    if _should_try_rapidocr_fallback(primary):
        rapid = run_rapidocr_subprocess(path)
        rapid_status = (
            OCR_STATUS_SUCCEEDED
            if rapid.get("ok") and (rapid.get("text") or "").strip()
            else (
                OCR_STATUS_EMPTY
                if rapid.get("ok") is False and not rapid.get("error_code")
                else (
                    OCR_STATUS_UNAVAILABLE
                    if rapid.get("error_code")
                    else OCR_STATUS_EMPTY
                )
            )
        )
        if rapid.get("ok") and (rapid.get("text") or "").strip():
            rapid_status = OCR_STATUS_SUCCEEDED
        elif rapid.get("error_code"):
            rapid_status = OCR_STATUS_UNAVAILABLE
        else:
            rapid_status = OCR_STATUS_EMPTY

        rapid_assessment = assess_ocr_text(
            str(rapid.get("text") or ""),
            runtime_status=rapid_status if rapid_status != OCR_STATUS_UNAVAILABLE else OCR_STATUS_FAILED,
            mean_confidence=rapid.get("mean_confidence"),
            provider=RAPID_PROVIDER,
            provider_version=str(rapid.get("model") or RAPID_PROVIDER),
        )
        # If subprocess hard-failed, keep assessment non-searchable.
        if rapid.get("error_code") and not (rapid.get("text") or "").strip():
            rapid_assessment = assess_ocr_text(
                "",
                runtime_status=OCR_STATUS_UNAVAILABLE,
                provider=RAPID_PROVIDER,
                provider_version=str(rapid.get("model") or RAPID_PROVIDER),
            )

        attempts.append(
            _safe_attempt_meta(
                provider=RAPID_PROVIDER,
                model=str(rapid.get("model") or ""),
                quality=str(rapid_assessment.ocr_quality_status),
                elapsed_ms=float(rapid.get("elapsed_ms") or 0),
                error_code=rapid.get("error_code"),
                mean_confidence=rapid.get("mean_confidence"),
            )
        )
        rapid_row = {
            "text": str(rapid.get("text") or ""),
            "provider": RAPID_PROVIDER,
            "model": rapid.get("model"),
            "status": rapid_status,
            "error_code": rapid.get("error_code"),
            "assessment": rapid_assessment,
            "elapsed_ms": float(rapid.get("elapsed_ms") or 0),
            "mean_confidence": rapid.get("mean_confidence"),
        }
        if rapid_row["text"].strip() or _quality_rank(rapid_assessment) > _quality_rank(primary["assessment"]):
            chosen = _pick_better_ocr(primary, rapid_row)
            fallback_used = chosen is rapid_row or chosen["provider"] == RAPID_PROVIDER

    assessment: OcrQualityAssessment = chosen["assessment"]
    text = str(chosen.get("text") or "")
    status = str(chosen.get("status") or OCR_STATUS_FAILED)
    if text and assessment.search_eligible:
        status = OCR_STATUS_SUCCEEDED
    elif not text and status == OCR_STATUS_SUCCEEDED:
        status = OCR_STATUS_EMPTY

    error_code = chosen.get("error_code")
    if status in {OCR_STATUS_UNAVAILABLE, OCR_STATUS_FAILED} and not error_code:
        error_code = OCR_FAILED

    safe_meta: Dict[str, Any] = {
        "ocr_provider": chosen["provider"],
        "ocr_model": chosen.get("model"),
        "ocr_fallback_used": bool(fallback_used),
        "ocr_mean_confidence": assessment.ocr_mean_confidence,
        "ocr_elapsed_ms": round(float(chosen.get("elapsed_ms") or 0), 2),
        "ocr_attempts": attempts,
    }
    if norm_meta:
        for key in ("source_format", "frame_count", "normalized_format", "selected_frame"):
            if key in norm_meta:
                safe_meta[key] = norm_meta[key]
    if error_code and status in {OCR_STATUS_UNAVAILABLE, OCR_STATUS_FAILED, OCR_STATUS_DISABLED}:
        safe_meta["ocr_error_code"] = str(error_code)

    return text, str(chosen["provider"]), status, (str(error_code) if error_code else None), safe_meta


def extract_image_chunks(path: Path, mime_type: str, original_filename: str | None = None) -> list[dict]:
    """Extract image caption plus optional OCR through the configured provider."""
    provider_name = normalize_provider()
    pipeline_version = str(
        _settings().attachment_ocr_pipeline_version or "quality-gated-rapidocr-v1"
    ).strip()
    if provider_name in {"", "tesseract"}:
        provider_name = "pytesseract"
    if provider_name in {"none", "off"}:
        provider_name = "disabled"

    metadata: dict = {
        "extraction_method": "image_metadata",
        "mime_type": mime_type,
        "ocr_pipeline_version": pipeline_version,
        "ocr_status": OCR_STATUS_UNAVAILABLE,
        "ocr_provider": provider_name,
        "ocr_language": _ocr_language_label(),
    }
    filename = original_filename or path.name
    caption_parts = [f"图片附件：{filename}", f"类型：{mime_type}"]

    try:
        from PIL import Image  # type: ignore

        with Image.open(path) as image:
            width, height = image.size
            metadata["width"] = width
            metadata["height"] = height
            caption_parts.append(f"尺寸：{width}x{height}")
    except Exception:
        # Do not persist raw probe exceptions/paths.
        pass

    chunks = [{
        "modality": "caption",
        "page_no": None,
        "slide_no": None,
        "content": "；".join(caption_parts),
        "metadata_json": metadata,
    }]

    ocr_text, provider, status, error_code, safe_meta = _run_ocr(path)
    chunks[0]["metadata_json"]["ocr_provider"] = provider
    chunks[0]["metadata_json"]["ocr_status"] = status
    chunks[0]["metadata_json"]["ocr_model"] = safe_meta.get("ocr_model")
    chunks[0]["metadata_json"]["ocr_fallback_used"] = bool(safe_meta.get("ocr_fallback_used"))
    chunks[0]["metadata_json"]["ocr_elapsed_ms"] = safe_meta.get("ocr_elapsed_ms")
    chunks[0]["metadata_json"]["ocr_attempts"] = safe_meta.get("ocr_attempts")
    for key in ("source_format", "frame_count", "normalized_format", "selected_frame"):
        if key in safe_meta:
            chunks[0]["metadata_json"][key] = safe_meta[key]
    if error_code and status in {OCR_STATUS_UNAVAILABLE, OCR_STATUS_FAILED, OCR_STATUS_DISABLED}:
        chunks[0]["metadata_json"]["ocr_error_code"] = error_code

    required = bool(_settings().attachment_image_ocr_required)

    quality = assess_ocr_text(
        ocr_text or "",
        runtime_status=status,
        mean_confidence=safe_meta.get("ocr_mean_confidence"),
        provider=provider,
        provider_version=str(safe_meta.get("ocr_model") or provider),
    )
    chunks[0]["metadata_json"] = merge_quality_into_metadata(
        chunks[0]["metadata_json"],
        quality,
    )

    # required: final search_eligible=false must fail even if Tesseract status=succeeded.
    if required and (
        status in {OCR_STATUS_UNAVAILABLE, OCR_STATUS_FAILED}
        or not quality.search_eligible
    ):
        raise OcrRequiredError(error_code=error_code or OCR_FAILED)

    # Caption-only retained; unusable OCR must not claim "recognized".
    if status == OCR_STATUS_SUCCEEDED and ocr_text and quality.search_eligible:
        lang = _ocr_language_label()
        for chunk in chunk_text(ocr_text):
            chunks.append({
                "modality": "ocr",
                "page_no": None,
                "slide_no": None,
                "content": chunk,
                "metadata_json": merge_quality_into_metadata(
                    {
                        "extraction_method": "ocr",
                        "mime_type": mime_type,
                        "ocr_pipeline_version": pipeline_version,
                        "ocr_provider": provider,
                        "ocr_model": safe_meta.get("ocr_model"),
                        "ocr_fallback_used": bool(safe_meta.get("ocr_fallback_used")),
                        "ocr_elapsed_ms": safe_meta.get("ocr_elapsed_ms"),
                        "ocr_language": lang,
                        "ocr_status": OCR_STATUS_SUCCEEDED,
                    },
                    quality,
                ),
            })
    else:
        # optional: caption-only indexed; explicitly not searchable.
        chunks[0]["metadata_json"]["ocr_text_searchable"] = False
        if not quality.search_eligible:
            chunks[0]["metadata_json"]["ocr_skipped_reason"] = (
                quality.quality_reason or "not_search_eligible"
            )

    return chunks


def extract_attachment_chunks(path: Path, mime_type: str, original_filename: str | None = None) -> list[dict]:
    """Extract searchable text chunks for supported attachment types."""
    if mime_type == "application/pdf":
        return extract_pdf_chunks(path)
    if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return extract_pptx_chunks(path)
    if mime_type.startswith("image/"):
        return extract_image_chunks(path, mime_type, original_filename=original_filename)
    raise RuntimeError(f"不支持的附件解析类型: {mime_type}")
